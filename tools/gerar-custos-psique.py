"""Planilha gerencial Psique: ausencia nunca vira custo zero. Requer openpyxl."""
from pathlib import Path
import ast, math, re, sys, zipfile
import xml.etree.ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from openpyxl.formula import Tokenizer

ROOT=Path(__file__).resolve().parents[1]
OUTPUT=ROOT/"FROID_Psique_Custos_Creditos_Atendimento.xlsx"
PENDING="Sem Capacidade de Apuração"
BRL='"R$" #,##0.00;[Red]-"R$" #,##0.00;"R$" 0.00'
NUM='#,##0.0000'
PCT='0.00%'

def sheet(wb,name,title,subtitle,headers,widths):
    ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False
    for row,text in [(1,title),(2,subtitle)]:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=len(headers))
        c=ws.cell(row,1,text)
        c.fill=PatternFill("solid",fgColor="142B49" if row==1 else "EDF2F7")
        c.font=Font(name="Calibri",size=16 if row==1 else 10,bold=row==1,color="FFFFFF" if row==1 else "142B49")
        c.alignment=Alignment(wrap_text=True,vertical="center");ws.row_dimensions[row].height=42
    for col,(h,width) in enumerate(zip(headers,widths),1):
        c=ws.cell(5,col,h);c.fill=PatternFill("solid",fgColor="142B49")
        c.font=Font(bold=True,color="FFFFFF");c.alignment=Alignment(wrap_text=True)
        ws.column_dimensions[c.column_letter].width=width
    ws.row_dimensions[5].height=32;ws.freeze_panes="B6"
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation="landscape";ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows="1:5"
    ws.oddFooter.center.text="FROID Psique | &A | &P / &N"
    return ws

def put(ws,ref,value=None,fmt=None,edit=False):
    c=ws[ref];c.value=value;c.font=Font(name="Calibri",size=11,color="142B49")
    c.alignment=Alignment(wrap_text=True,vertical="center")
    c.fill=PatternFill("solid",fgColor="FFF2CC" if edit else "E2F0D9" if isinstance(value,str) and value.startswith("=") else "DBEAFE")
    if fmt:c.number_format=fmt
    ws.row_dimensions[c.row].height=max(ws.row_dimensions[c.row].height or 15,44)
    return c

def dv(ws,cells,low=0,high=1e12,values=None):
    d=DataValidation(type="list" if values else "decimal",operator=None if values else "between",formula1='"'+",".join(values)+'"' if values else str(low),formula2=None if values else str(high),allow_blank=True)
    d.errorTitle="Entrada inválida";d.error="Informe valor válido; deixe vazio se não apurado."
    d.showErrorMessage=True;d.errorStyle="stop";ws.add_data_validation(d);d.add(cells)

def total(ref,n):
    return f'=IF(COUNT({ref})={n},SUM({ref}),"{PENDING}")'

def reference_packages():
    tree=ast.parse((ROOT/"froid-server/subscriptions.py").read_text(encoding="utf-8"))
    constants={n.target.id:ast.literal_eval(n.value) for n in tree.body
               if isinstance(n,ast.AnnAssign) and isinstance(n.target,ast.Name)
               and n.target.id in {"SESSION_PACKAGES","SUBSCRIPTION_PLANS"}}
    page=(ROOT/"froid-site/precos.html").read_text(encoding="utf-8")
    rows=[]
    for plan,body in re.findall(r"<h3>FROID (PRO|PLUS|MASTER)</h3>(.*?)(?=<h3>FROID|</section>)",page,re.S):
        code=plan.lower()
        packages=[p for p in constants["SESSION_PACKAGES"].values() if p["plan_code"]==code]
        for count,money in re.findall(r'<span class="qtd">(\d+) créditos de atendimento</span><span class="valor">R\$ ([\d.,]+)',body):
            count=int(count)
            candidates=[p for p in packages if p["sessions"]==count]
            if not candidates and len(packages)==1:candidates=packages
            assert len(candidates)==1, "Catálogo alterado: revisar correspondência do pacote"
            p=candidates[0]
            rows.append((plan,count,constants["SUBSCRIPTION_PLANS"][code]["entitlements"]["organization_members"],
                         float(money.replace(".","").replace(",",".")),p["sessions"],p["prices"]["brl"]["total_amount_minor"]/100))
    assert len(rows)==5,"Revisar catálogo de pacotes: não foram encontrados os cinco pacotes de referência."
    return rows

def build():
    wb=Workbook();wb.remove(wb.active)
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
    ws=sheet(wb,"Guia","FROID PSIQUE | Custo por crédito de atendimento","24/09/2026 | Estudo gerencial. Não aprova nem altera preços ou configurações Stripe.",["Etapa","Como utilizar"],[30,115])
    guide=[
    ("Unidade","Um crédito de atendimento corresponde a um atendimento na regra atual. Sessão continua sendo o registro clínico; duração e regra de consumo não foram alteradas."),
    ("Preenchimento","Amarelo: entrada. Azul: referência. Verde: fórmula. Vazio = não apurado. Zero somente quando houver confirmação de custo/consumo nulo."),
    ("Ordem","Preencha Premissas, Fixos, Variáveis e IA; consulte Resultado, Pacotes e Cenários."),
    ("Exclusões","Não se aplica exige justificativa. Itens pendentes bloqueiam custo completo e preço mínimo, evitando que um subtotal pareça o total."),
    ("Regime","Modelo mensal de equilíbrio: receitas de créditos pagos utilizados, uso integral dos pacotes e taxas de compra rateadas. Não é fluxo de caixa de vendas antecipadas."),
    ("Gratuitos","Atendimentos gratuitos também consomem APIs e estrutura. O custo deles é suportado pelos créditos pagos. Desconto de pioneiro encerrado; cortesia não foi cancelada."),
    ("Não duplicar","Cada gasto aparece uma vez. APIs são calculadas em IA. Suporte pode ser folha fixa OU variável por atendimento; não repetir horas nem faturas."),
    ("Rateio","Infra compartilhada com NR-1, site e desenvolvimento exige critério documentado. Não atribuir tudo ao Psique sem apuração."),
    ("APIs","Use minutos efetivamente enviados de paciente e profissional, incluindo reenvios. Tokens incluem prompts, histórico, resumos, análises, relatórios, consultas e fallback. max_tokens não é consumo."),
    ("Histórico","Armazenamento e backup incluem o acervo acumulado. Use faturas ou projeções documentadas; não presumir exclusão de registros após acabar o crédito."),
    ("Câmbio e impostos","Informe câmbio efetivo com IOF/spread. A contabilidade deve apurar a alíquota de tributos. Nenhum regime tributário foi presumido."),
    ("Pagamentos","Referências contém taxas públicas. Em Premissas, informe a tarifa realmente contratada ou ponderada pelo mix. Tarifa fixa incide por pacote, não por crédito."),
    ("Pendências","Sem faturas e consumo não há custo final medido. O total fica como Sem Capacidade de Apuração. Tarifas de referência não são gastos da empresa."),
    ("Preços e MASTER","Pacotes registra site e catálogo consultados. A divergência MASTER permanece visível; valor de teste não foi adotado como recomendação."),
    ("Cenários","Volumes e preços ficam vazios. Hipóteses preenchidas por você geram projeções, não medições. Mudança de capacidade exige novo orçamento."),
    ("Escopo","Inclui infraestrutura, APIs, armazenamento, pessoal, pró-labore, suporte, desenvolvimento, aquisição, jurídico, contabilidade, seguros, tributos, taxas e perdas."),
    ("Fonte","Modelos em produção conferidos em 24/09/2026: gpt-4o-mini, gpt-4o-transcribe, gemini-1.5-pro. Não houve acesso a extratos de fornecedores nem a dados clínicos."),
    ]
    for r,(a,b) in enumerate(guide,6):
        put(ws,f"A{r}",a);put(ws,f"B{r}",b);ws.row_dimensions[r].height=58
    ws=sheet(wb,"Premissas","ENTRADAS | Volume, preço e cobrança","Valores confirmados ou hipóteses identificadas na fonte. Digite percentuais com %. Custos ausentes não são zero.",["Premissa","Valor","Unidade","Fonte / observação"],[48,32,20,100])
    rows={
    6:("Créditos pagos utilizados no mês","créditos","Base de rateio: maior que zero."),
    7:("Créditos gratuitos utilizados no mês","créditos","Zero somente se não houver uso gratuito."),
    8:("Atendimentos totais no mês","atendimentos","Pagos + gratuitos."),
    9:("Duração média observada","minutos","Contexto: não substitui minutos faturados de transcrição."),
    10:("Créditos por pacote de referência","créditos/pacote","Para mix, média comprovada de créditos por pagamento."),
    11:("Preço a avaliar por crédito","BRL/crédito","Sua decisão comercial, ainda não aprovada/publicada."),
    12:("Margem operacional desejada","% receita","Margem sobre receita, não markup."),
    13:("USD para BRL efetivo","BRL/USD","Inclua spread e encargos cambiais; não duplicar em outra linha."),
    14:("EUR para BRL efetivo","BRL/EUR","Necessário apenas para faturas em EUR."),
    15:("Tributos efetivos sobre faturamento","% receita","Apurar com a contabilidade."),
    16:("Comissões variáveis","% receita","Não duplicar valores lançados em Fixos/Variáveis."),
    17:("Perdas e reembolsos líquidos","% receita","Provisão informada; tarifas de contestação ficam em Fixos."),
    18:("Tarifa percentual do pagamento","% receita","Contratada ou média ponderada: consulte Referências."),
    19:("Tarifa fixa por pagamento","BRL/transação","Por pacote comprado, não por crédito."),
    20:("Outras tarifas percentuais","% receita","Parcelamento, Billing, conversão ou adicionais não incluídos acima."),
    21:("Desconto de pioneiro vigente","%","0% por determinação do proprietário em 24/09/2026."),
    }
    for r,(label,unit,note) in rows.items():
        put(ws,f"A{r}",label);put(ws,f"C{r}",unit);put(ws,f"D{r}",note)
        fmt=PCT if r in [12,15,16,17,18,20,21] else BRL if r in [11,19] else NUM
        put(ws,f"B{r}",0 if r==21 else None,fmt,edit=r not in [8,21])
        if r not in [8,21]:dv(ws,f"B{r}",high=.999999 if fmt==PCT else 1e12)
    put(ws,"B8",total("B6:B7",2),"0")
    ws=sheet(wb,"Referencias","REFERÊNCIAS | Tarifas públicas","Consulta em 24/09/2026. Não substitui contrato/fatura. Gemini configurado não localizado na tabela atual: tarifa pendente.",["Serviço","Tarifa","Unidade","Fonte","Observação"],[43,24,25,78,90])
    ai="https://developers.openai.com/api/docs/"
    stripe="https://stripe.com/br/pricing"
    refs=[
    ("gpt-4o-transcribe",.006,"USD/minuto",ai+"pricing","Estimativa oficial por minuto; conciliar com usage/fatura."),
    ("gpt-4o-mini entrada",.15,"USD/1M tokens",ai+"models/gpt-4o-mini","Tarifa padrão sem cache; não usar preço Batch."),
    ("gpt-4o-mini cache",.075,"USD/1M tokens",ai+"models/gpt-4o-mini","Só tokens cobrados com cache."),
    ("gpt-4o-mini saída",.60,"USD/1M tokens",ai+"models/gpt-4o-mini","Uso medido, não limite de geração."),
    ("text-embedding-3-small",.02,"USD/1M tokens",ai+"models/text-embedding-3-small","Confirmar consultas e reindexações atribuíveis ao Psique."),
    ("Gemini entrada (gemini-1.5-pro)",None,"USD/1M tokens","https://ai.google.dev/gemini-api/docs/pricing","Tarifa atual não localizada para modelo configurado. Confirmar uso/modelo/fatura."),
    ("Gemini saída (gemini-1.5-pro)",None,"USD/1M tokens","https://ai.google.dev/gemini-api/docs/pricing","Ausência não significa custo zero; pode haver fallback OpenAI."),
    ("Stripe cartão nacional — percentual",.0399,"% transação",stripe,"Confirmar tarifa contratada."),
    ("Stripe cartão nacional — fixa",.39,"BRL/transação",stripe,"Distribuir pelos créditos comprados."),
    ("Stripe internacional — adicional",.02,"% transação",stripe,"Adicional à tarifa nacional; verificar condições de conversão."),
    ("Stripe Pix",.0119,"% transação",stripe,"Por convite segundo página consultada; não assume habilitação."),
    ("Stripe boleto",3.45,"BRL/boleto pago",stripe,"Alternativa; não somar ao cartão na mesma venda."),
    ("Stripe contestação recebida",55,"BRL/evento",stripe,"Não presumir ocorrência."),
    ("Stripe refutação manual",55,"BRL/evento",stripe,"Tarifa devolvida em disputas ganhas; registrar custo líquido."),
    ("Catálogo comercial","Confirmado","código","froid-server/subscriptions.py","MASTER diverge do site; ver Pacotes."),
    ("Infra compartilhada","Ratear","código","docker-compose.yml","Não alocar NR-1 e desenvolvimento integralmente ao Psique."),
    ]
    for r,values in enumerate(refs,6):
        for col,v in enumerate(values,1):
            put(ws,f"{chr(64+col)}{r}",v,PCT if col==2 and r in [13,15,16] else NUM if col==2 and isinstance(v,(float,int)) else None,edit=col==2 and r in [11,12])
        if values[3].startswith("https"):ws[f"D{r}"].hyperlink=values[3]
    dv(ws,"B11:B12")
    ws=sheet(wb,"IA","APIs | Custo por atendimento","Quantidade média por atendimento, incluindo gratuitos. Some todo uso pago; entrada com cache não pode ser contada novamente sem cache.",["Componente","Quantidade","Unidade","Tarifa USD","Divisor","USD/atendimento","BRL/atendimento","Fonte da medição"],[47,22,23,26,14,32,33,80])
    ia=[("Transcrição paciente","minutos",6,1),("Transcrição profissional","minutos",6,1),
    ("Texto OpenAI — entrada sem cache","tokens",7,1000000),("Texto OpenAI — entrada cache","tokens",8,1000000),
    ("Texto OpenAI — saída","tokens",9,1000000),("Embeddings Psique","tokens",10,1000000),
    ("Gemini — entrada","tokens",11,1000000),("Gemini — saída","tokens",12,1000000),
    ("Outras APIs / conciliação comprovada","unidades",None,1)]
    for r,(label,unit,ref,divisor) in enumerate(ia,6):
        put(ws,f"A{r}",label);put(ws,f"B{r}",edit=True,fmt=NUM);put(ws,f"C{r}",unit)
        put(ws,f"D{r}",f'=IF(ISNUMBER(Referencias!B{ref}),Referencias!B{ref},"{PENDING}")' if ref else None,NUM,edit=not ref)
        put(ws,f"E{r}",divisor);put(ws,f"H{r}",edit=True)
        put(ws,f"F{r}",f'=IF(COUNT(B{r})<>1,"{PENDING}",IF(B{r}=0,0,IF(COUNT(D{r})=1,B{r}*D{r}/E{r},"{PENDING}")))',NUM)
        put(ws,f"G{r}",f'=IF(ISNUMBER(F{r}),IF(F{r}=0,0,IF(AND(ISNUMBER(Premissas!B13),Premissas!B13>0),F{r}*Premissas!B13,"{PENDING}")),"{PENDING}")',BRL)
    dv(ws,"B6:B14");dv(ws,"D14")
    put(ws,"A16","TOTAL IA POR ATENDIMENTO");put(ws,"G16",total("G6:G14",9),BRL)
    put(ws,"A18","PENDÊNCIAS");put(ws,"G18","=9-COUNT(G6:G14)","0")
    ws=sheet(wb,"Fixos","CUSTOS MENSAIS | Faturas e estrutura","Classifique todas as linhas. Não se aplica exige justificativa. Valor 0 exige confirmação de ausência de gasto. Documente o rateio.",["Item","Tratamento","Valor da fatura","Moeda","Rateio Psique","BRL/mês","Fonte / justificativa","Situação"],[47,24,24,14,20,34,92,18])
    costs=["Servidor Hetzner / computação","Banco de dados / volume adicional","Backups / snapshots / recuperação",
    "Armazenamento de históricos e documentos","TURN / conectividade / tráfego contratado","CDN / DNS / proteção de borda",
    "Domínio (mensalizar custo anual)","E-mail transacional / SMTP","Monitoramento / logs / alertas","Segurança / auditoria / testes",
    "Suporte fixo / atendimento ao cliente","Manutenção / desenvolvimento recorrente","Pró-labore e encargos dos sócios",
    "Equipe administrativa / encargos","Equipe técnica / encargos","Equipe comercial / parcela fixa","Contabilidade / obrigações fiscais",
    "Jurídico / contratos / LGPD / encarregado","Seguros / responsabilidade / cibernético","Marketing / mídia / aquisição",
    "Ferramentas de trabalho e licenças","Aluguel / internet / energia rateados","Tarifas bancárias fixas",
    "Depreciação / amortização de equipamentos","Amortização do desenvolvimento inicial","Treinamento / implantação de clientes",
    "Contestações Stripe — tarifas líquidas","Outros custos fixos ou periódicos"]
    for r,label in enumerate(costs,6):
        put(ws,f"A{r}",label);put(ws,f"B{r}","Pendente",edit=True)
        for col in ["C","D","E","G"]:put(ws,f"{col}{r}",edit=True,fmt=PCT if col=="E" else NUM if col=="C" else None)
        conv=f'IF(D{r}="BRL",1,IF(D{r}="USD",IF(AND(ISNUMBER(Premissas!B13),Premissas!B13>0),Premissas!B13,"{PENDING}"),IF(D{r}="EUR",IF(AND(ISNUMBER(Premissas!B14),Premissas!B14>0),Premissas!B14,"{PENDING}"),"{PENDING}")))'
        f=f'=IF(B{r}="Não se aplica",IF(G{r}<>"",0,"{PENDING}"),IF(AND(B{r}="Informado",COUNT(C{r},E{r})=2,C{r}>=0,E{r}>=0,E{r}<=1,G{r}<>""),IF(ISNUMBER({conv}),C{r}*E{r}*{conv},"{PENDING}"),"{PENDING}"))'
        put(ws,f"F{r}",f,BRL);put(ws,f"H{r}",f'=IF(ISNUMBER(F{r}),"OK","Pendente")')
    dv(ws,"B6:B33",values=["Pendente","Informado","Não se aplica"]);dv(ws,"D6:D33",values=["BRL","USD","EUR"])
    dv(ws,"C6:C33");dv(ws,"E6:E33",high=1)
    put(ws,"A35","TOTAL MENSAL PSIQUE");put(ws,"F35",total("F6:F33",28),BRL)
    put(ws,"A37","PENDÊNCIAS");put(ws,"F37","=28-COUNT(F6:F33)","0");ws.auto_filter.ref="A5:H33"
    ws=sheet(wb,"Variaveis","OUTROS CUSTOS | Por atendimento","Somente incrementais fora de IA e Fixos. Faturas do acervo total ficam em Fixos. Não duplicar parcelas já incluídas na mensalidade.",["Item","Tratamento","Qtd/atendimento","BRL/unidade","BRL/atendimento","Fonte / justificativa"],[48,24,24,24,34,95])
    items=["Tráfego excedente / relay TURN (GB)","Armazenamento incremental não faturado em Fixos","Mensagens / e-mails / notificações",
    "Suporte variável (minutos)","Implantação/treinamento amortizado por crédito","Serviços terceirizados por atendimento",
    "Emissão de nota / documento (parcela variável)","Outras APIs em BRL não lançadas em IA","Outros custos incrementais"]
    for r,label in enumerate(items,6):
        put(ws,f"A{r}",label);put(ws,f"B{r}","Pendente",edit=True)
        for c in ["C","D","F"]:put(ws,f"{c}{r}",edit=True,fmt=NUM if c!="F" else None)
        put(ws,f"E{r}",f'=IF(B{r}="Não se aplica",IF(F{r}<>"",0,"{PENDING}"),IF(AND(B{r}="Informado",COUNT(C{r},D{r})=2,C{r}>=0,D{r}>=0,F{r}<>""),C{r}*D{r},"{PENDING}"))',BRL)
    dv(ws,"B6:B14",values=["Pendente","Informado","Não se aplica"]);dv(ws,"C6:D14")
    put(ws,"A16","TOTAL POR ATENDIMENTO");put(ws,"E16",total("E6:E14",9),BRL)
    put(ws,"A18","PENDÊNCIAS");put(ws,"E18","=9-COUNT(E6:E14)","0");ws.auto_filter.ref="A5:F14"
    ws=sheet(wb,"Resultado","RESULTADO | Custo, margem e equilíbrio","Custo completo só aparece quando todas as dependências forem preenchidas. Projeções não são faturas. As premissas definem a precisão.",["Indicador","Valor","Interpretação"],[56,38,110])
    lines={
    6:("Custo fixo mensal",'=Fixos!F35',"Parcela atribuída ao Psique."),
    7:("IA por atendimento",'=IA!G16',"Tarifa pública × consumo informado × câmbio efetivo."),
    8:("Outros variáveis por atendimento",'=Variaveis!E16',"Incrementais fora de IA."),
    9:("Variável total por atendimento",total("B7:B8",2),"Pago ou gratuito."),
    10:("Custos operacionais totais no mês",f'=IF(AND(COUNT(B6,B9,Premissas!B8)=3,Premissas!B8>=0),B6+B9*Premissas!B8,"{PENDING}")',"Fixo + variável × atendimentos pagos e gratuitos."),
    11:("Custo operacional por crédito pago",f'=IF(AND(COUNT(B10,Premissas!B6)=2,Premissas!B6>0),B10/Premissas!B6,"{PENDING}")',"Inclui gratuitos e estrutura. Ainda sem tributos/taxas da venda."),
    12:("Deduções percentuais sobre receita",total("Premissas!B15:B18,Premissas!B20",5),"Tributos + comissões + perdas + pagamento + adicionais."),
    13:("Tarifa fixa por crédito pago",f'=IF(AND(COUNT(Premissas!B19,Premissas!B10)=2,Premissas!B10>0),Premissas!B19/Premissas!B10,"{PENDING}")',"Tarifa por transação dividida pelos créditos do pacote."),
    14:("Preço de equilíbrio por crédito",f'=IF(COUNT(B11:B13)=3,IF(B12<1,(B11+B13)/(1-B12),"Inviável"),"{PENDING}")',"Cobre custos informados com margem zero."),
    15:("Preço mínimo para margem desejada",f'=IF(COUNT(B11:B13,Premissas!B12)=4,IF(B12+Premissas!B12<1,(B11+B13)/(1-B12-Premissas!B12),"Inviável"),"{PENDING}")',"Custo dividido pela receita restante após taxas e margem-alvo."),
    17:("Receita mensal ao preço avaliado",f'=IF(AND(COUNT(Premissas!B6,Premissas!B11)=2,Premissas!B6>0,Premissas!B11>0),Premissas!B6*Premissas!B11,"{PENDING}")',"Receita gerencial dos créditos pagos utilizados; não caixa das compras antecipadas."),
    18:("Deduções percentuais em BRL",f'=IF(COUNT(B17,B12)=2,B17*B12,"{PENDING}")',"Não inclui tarifa fixa."),
    19:("Tarifas fixas de pagamento no mês",f'=IF(COUNT(B13,Premissas!B6)=2,B13*Premissas!B6,"{PENDING}")',"Rateio pelos créditos comprados."),
    20:("Custos totais com cobrança e tributos",total("B10,B18:B19",3),"Operação + deduções + tarifas fixas."),
    21:("Resultado operacional mensal",f'=IF(COUNT(B17,B20)=2,B17-B20,"{PENDING}")',"Receita menos custos informados."),
    22:("Margem operacional",f'=IF(COUNT(B17,B21)=2,IF(B17>0,B21/B17,"Inviável"),"{PENDING}")',"Projeção se as entradas forem hipóteses."),
    23:("Créditos pagos para atingir equilíbrio",f'=IF(COUNT(B6,B9,B12,B13,Premissas!B7,Premissas!B11)=6,IF(Premissas!B11*(1-B12)-B9-B13>0,ROUNDUP((B6+B9*Premissas!B7)/(Premissas!B11*(1-B12)-B9-B13),0),"Inviável"),"{PENDING}")',"Mantém estrutura, gratuitos e custo unitário constantes. Mudança de capacidade exige recalcular."),
    25:("Itens de custo ainda pendentes","=SUM(IA!G18,Fixos!F37,Variaveis!E18)","Além destes, premissas comerciais/câmbio também podem faltar."),
    }
    for r,(a,f,note) in lines.items():
        put(ws,f"A{r}",a);put(ws,f"B{r}",f,PCT if r in [12,22] else "0" if r in [23,25] else BRL);put(ws,f"C{r}",note)
    ws=sheet(wb,"Pacotes","PACOTES | Referências e valores em avaliação","Nenhum preço novo é aprovado aqui. MASTER: site e catálogo divergem; não usar valor de teste como decisão comercial.",["Plano","Créditos site","Usuários","Total site BRL","Créditos catálogo","Total catálogo BRL","Preço/crédito site","Novo total (entrada)","Novo preço/crédito","Custo operacional","Taxas/tributos","Resultado pacote","Margem","Mínimo p/ margem-alvo"],[18,18,15,22,20,22,23,25,25,33,33,33,24,34])
    pkgs=reference_packages()
    for r,vals in enumerate(pkgs,6):
        for c,v in enumerate(vals,1):put(ws,f"{chr(64+c)}{r}",v,BRL if c in [4,6] else None)
        put(ws,f"G{r}",f"=D{r}/B{r}",BRL);put(ws,f"H{r}",edit=True,fmt=BRL)
        put(ws,f"I{r}",f'=IF(COUNT(H{r})=1,H{r}/B{r},"{PENDING}")',BRL)
        put(ws,f"J{r}",f'=IF(ISNUMBER(Resultado!B11),Resultado!B11*B{r},"{PENDING}")',BRL)
        put(ws,f"K{r}",f'=IF(COUNT(H{r},Resultado!B12,Premissas!B19)=3,H{r}*Resultado!B12+Premissas!B19,"{PENDING}")',BRL)
        put(ws,f"L{r}",f'=IF(COUNT(H{r},J{r},K{r})=3,H{r}-J{r}-K{r},"{PENDING}")',BRL)
        put(ws,f"M{r}",f'=IF(COUNT(H{r},L{r})=2,IF(H{r}>0,L{r}/H{r},"Inviável"),"{PENDING}")',PCT)
        put(ws,f"N{r}",f'=IF(COUNT(J{r},Resultado!B12,Premissas!B12,Premissas!B19)=4,IF(Resultado!B12+Premissas!B12<1,(J{r}+Premissas!B19)/(1-Resultado!B12-Premissas!B12),"Inviável"),"{PENDING}")',BRL)
    dv(ws,"H6:H10",low=.01)
    for r,text in [
    (12,"MASTER: site = 200 créditos / R$ 4.888; catálogo consultado = 25 / R$ 20. Divergência registrada, sem adoção como novo preço."),
    (14,"Preço por crédito = total do pacote / quantidade. Não utiliza o campo unitário de referência comercial, que pode divergir dessa divisão."),
    (16,"Fontes: froid-server/subscriptions.py e https://www.froid.com.br/precos.html. Preços de referência, sujeitos à decisão do proprietário.")]:
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=14);put(ws,f"A{r}",text)
    ws=sheet(wb,"Cenarios","CENÁRIOS | Volume e preço escolhidos por você","Nenhum volume ou preço foi presumido. Mesmos fixos, variável unitário e taxas do cenário base. Capacidade extra deve ser orçada.",["Cenário","Pagos/mês","Gratuitos/mês","BRL/crédito","Receita mensal","Custo completo","Resultado mensal","Margem","Preço equilíbrio"],[22,22,24,27,33,34,34,24,34])
    for r in range(6,11):
        put(ws,f"A{r}",f"Cenário {r-5}")
        for c in ["B","C","D"]:put(ws,f"{c}{r}",edit=True,fmt=BRL if c=="D" else "0")
        put(ws,f"E{r}",f'=IF(AND(COUNT(B{r},D{r})=2,B{r}>0,D{r}>0),B{r}*D{r},"{PENDING}")',BRL)
        put(ws,f"F{r}",f'=IF(AND(COUNT(B{r}:D{r},E{r},Resultado!B6,Resultado!B9,Resultado!B12,Resultado!B13)=8,B{r}>0,C{r}>=0),Resultado!B6+Resultado!B9*(B{r}+C{r})+E{r}*Resultado!B12+Resultado!B13*B{r},"{PENDING}")',BRL)
        put(ws,f"G{r}",f'=IF(COUNT(E{r}:F{r})=2,E{r}-F{r},"{PENDING}")',BRL)
        put(ws,f"H{r}",f'=IF(COUNT(E{r},G{r})=2,IF(E{r}>0,G{r}/E{r},"Inviável"),"{PENDING}")',PCT)
        put(ws,f"I{r}",f'=IF(AND(COUNT(B{r}:C{r},Resultado!B6,Resultado!B9,Resultado!B12,Resultado!B13)=6,B{r}>0,C{r}>=0),IF(Resultado!B12<1,(Resultado!B6/B{r}+Resultado!B9*(1+C{r}/B{r})+Resultado!B13)/(1-Resultado!B12),"Inviável"),"{PENDING}")',BRL)
    dv(ws,"B6:C10");dv(ws,"D6:D10",low=.01)
    return wb

# Avalia o subconjunto de fórmulas utilizado; IF é lazy, como no Excel.
# Também grava cache verdadeiro para visualizadores que não recalculam.
def evaluate(wb):
    cache={};active=set()
    def cell(sheet,coord):
        key=(sheet,coord.replace("$",""))
        if key in cache:return cache[key]
        if key in active:raise ValueError("Referência circular: "+str(key))
        value=wb[sheet][key[1]].value
        if not isinstance(value,str) or not value.startswith("="):return value
        active.add(key);tokens=[t for t in Tokenizer(value).items if t.type!="WHITE-SPACE"];pos=0
        def expr(minp=0):
            nonlocal pos
            t=tokens[pos];pos+=1
            if t.type=="FUNC" and t.subtype=="OPEN":
                name=t.value[:-1];args=[]
                while tokens[pos].subtype!="CLOSE":
                    args.append(expr())
                    if tokens[pos].type=="SEP":pos+=1
                    else:break
                assert tokens[pos].subtype=="CLOSE";pos+=1;node=("fn",name,args)
            elif t.type=="PAREN" and t.subtype=="OPEN":
                node=expr();assert tokens[pos].subtype=="CLOSE";pos+=1
            elif t.type=="OPERATOR-PREFIX":node=("op",t.value,("num",0),expr(4))
            elif t.subtype=="NUMBER":node=("num",float(t.value))
            elif t.subtype=="TEXT":node=("str",t.value[1:-1].replace('""','"'))
            elif t.subtype=="RANGE":node=("ref",t.value)
            else:raise ValueError((value,t))
            precedence={"=":1,"<>":1,">":1,"<":1,">=":1,"<=":1,"+":2,"-":2,"*":3,"/":3}
            while pos<len(tokens) and tokens[pos].type=="OPERATOR-INFIX":
                op=tokens[pos].value;pr=precedence[op]
                if pr<minp:break
                pos+=1;node=("op",op,node,expr(pr+1))
            return node
        tree=expr();assert pos==len(tokens),(value,pos)
        def flat(values):
            for v in values:
                if isinstance(v,list):yield from flat(v)
                else:yield v
        def number(v):return isinstance(v,(int,float)) and not isinstance(v,bool)
        def ev(n):
            tag=n[0]
            if tag in ["num","str"]:return n[1]
            if tag=="ref":
                ref=n[1].replace("$","");sn,ref=ref.rsplit("!",1) if "!" in ref else (sheet,ref);sn=sn.strip("'")
                if ":" in ref:return [cell(sn,c.coordinate) for row in wb[sn][ref] for c in row]
                return cell(sn,ref)
            if tag=="op":
                a,b=ev(n[2]),ev(n[3]);op=n[1]
                if op in ["=","<>"]:
                    result=("" if a is None else a)==("" if b is None else b)
                    return result if op=="=" else not result
                a=0 if a is None else a;b=0 if b is None else b
                same=type(a)==type(b) or number(a) and number(b)
                if op==">":return a>b if same else isinstance(a,str)
                if op=="<":return a<b if same else not isinstance(a,str)
                if op==">=":return a>=b if same else isinstance(a,str)
                if op=="<=":return a<=b if same else not isinstance(a,str)
                if op=="+":return a+b
                if op=="-":return a-b
                if op=="*":return a*b
                if op=="/":return a/b
            name,args=n[1],n[2]
            if name=="IF":return ev(args[1]) if ev(args[0]) else ev(args[2])
            vals=list(flat([ev(a) for a in args]))
            if name=="COUNT":return sum(number(v) for v in vals)
            if name=="SUM":return sum(v for v in vals if number(v))
            if name=="ISNUMBER":return number(vals[0])
            if name=="AND":return all(vals)
            if name=="ROUNDUP":return math.ceil(vals[0]*10**vals[1])/10**vals[1]
            raise ValueError(name)
        result=ev(tree);active.remove(key);cache[key]=result;return result
    for ws in wb:
        for row in ws:
            for c in row:
                if c.data_type=="f":cell(ws.title,c.coordinate)
    return cache

def save(wb,path):
    values=evaluate(wb);wb.save(path)
    ns={"m":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as z:files={n:z.read(n) for n in z.namelist()}
    for i,ws in enumerate(wb,1):
        name=f"xl/worksheets/sheet{i}.xml";root=ET.fromstring(files[name])
        for c in root.findall(".//m:c",ns):
            if c.find("m:f",ns) is None:continue
            value=values[(ws.title,c.attrib["r"])];v=c.find("m:v",ns)
            if v is None:v=ET.SubElement(c,"{"+ns["m"]+"}v")
            if isinstance(value,str):c.set("t","str")
            else:c.attrib.pop("t",None)
            v.text=str(value)
        files[name]=ET.tostring(root,encoding="utf-8",xml_declaration=True)
    with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED) as z:
        for n,data in files.items():z.writestr(n,data)

def self_check():
    wb=build();v=evaluate(wb)
    assert v["Resultado","B15"]==PENDING
    assert v["Resultado","B25"]==46
    assert v["Pacotes","G6"]==19.8
    # Números fictícios APENAS em memória de teste, nunca na planilha entregue.
    p=wb["Premissas"]
    for r,x in {6:100,7:10,10:10,11:20,12:.2,13:5,14:6,15:.1,16:0,17:0,18:.04,19:.4,20:0}.items():p[f"B{r}"]=x
    for r in range(6,15):wb["IA"][f"B{r}"]=0
    wb["IA"]["B6"]=10
    for r in range(6,34):wb["Fixos"][f"B{r}"]="Não se aplica";wb["Fixos"][f"G{r}"]="Somente teste"
    for key,x in {"B6":"Informado","C6":100,"D6":"BRL","E6":1}.items():wb["Fixos"][key]=x
    for r in range(6,15):wb["Variaveis"][f"B{r}"]="Não se aplica";wb["Variaveis"][f"F{r}"]="Somente teste"
    v=evaluate(wb)
    assert math.isclose(v["Resultado","B9"],.3)
    assert math.isclose(v["Resultado","B10"],133)
    assert math.isclose(v["Resultado","B21"],1583)
    assert math.isclose(v["Resultado","B15"],1.37/.66)
    # Pacotes e cenários devem usar tarifa fixa por compra, não por crédito.
    wb["Pacotes"]["H6"]=200
    for key,x in {"B6":100,"C6":10,"D6":20}.items():wb["Cenarios"][key]=x
    v=evaluate(wb)
    assert math.isclose(v["Pacotes","K6"],28.4)
    assert math.isclose(v["Pacotes","L6"],158.3)
    assert math.isclose(v["Cenarios","G6"],1583)
    wb["Fixos"]["D6"]="EUR";wb["Fixos"]["E6"]=.5
    assert math.isclose(evaluate(wb)["Fixos","F6"],300)
    wb["Fixos"]["D6"]="BRL";wb["Fixos"]["E6"]=1
    p["B12"]=.9
    assert evaluate(wb)["Resultado","B15"]=="Inviável"
    p["B12"]=.2
    wb["Fixos"]["G7"]=None
    assert evaluate(wb)["Resultado","B15"]==PENDING
    wb["Fixos"]["G7"]="Somente teste";p["B6"]=0
    assert evaluate(wb)["Resultado","B15"]==PENDING
    print("OK: ausência, zero confirmado, exclusão documentada, cortesia, taxas e margem.")

if __name__=="__main__":
    self_check()
    dest=Path(sys.argv[1]) if len(sys.argv)>1 else OUTPUT
    save(build(),dest)
    reread=load_workbook(dest,data_only=True)
    assert reread["Resultado"]["B15"].value==PENDING
    assert reread["Premissas"]["B6"].value is None
    print(dest)
